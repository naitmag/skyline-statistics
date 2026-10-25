#!/usr/bin/env python3
"""
Обновление данных страницы index.html из xlsx-книги статистики.

Использование (xlsx можно указать и ссылкой на экспорт Google Таблицы:
https://docs.google.com/spreadsheets/d/<ID>/export?format=xlsx):
    pip install openpyxl
    python update_stats.py SKYLINE_Statistics.xlsx index.html            # обновить страницу (копия .bak)
    python update_stats.py SKYLINE_Statistics.xlsx index.html -o new.html # записать в другой файл
    python update_stats.py SKYLINE_Statistics.xlsx index.html --check    # только отчёт, без записи

Скрипт читает листы матчей («Игра №N …») и листы «Динамика», пересобирает блок
`const DATA = …` на странице и печатает отчёт сверки. Матч попадает на страницу,
если в строке «Очки команды» есть хотя бы одно очко.

Сверка (всё, что не сошлось, печатается как ПРЕДУПРЕЖДЕНИЕ):
  * очки партии = чистые очки (атака++, блок+, эйсы) + ошибки соперника;
  * ошибки игроков = ошибки команды;
  * лог «Динамика» (счёт по розыгрышам) против таблицы: очки, наши ошибки, чистые очки, ошибки соперника;
  * скачки счёта в логе, игроки в логе, которых нет в списке, пустые листы «Динамика»;
  * итог в книге против суммы по партиям (для «Связки» известное расхождение: формула итога не берёт «Связка −»).
"""
import argparse, datetime, json, os, re, shutil, sys, tempfile, urllib.request
import openpyxl

# ----- настройки ------------------------------------------------------------
# Как игроки названы в логе «Динамика» -> как в шапке листа матча.
NICKNAMES = {
    'Ваня': 'Иван Шумик', 'Женя': 'Евгений Мартопляс', 'Миша': 'Михаил Ярмолович',
    'Макс': 'Максим Битько', 'Даня': 'Даниил Тюрюханов', 'Артем(С)': 'Артем Бондарчук',
    'Аретем(С)': 'Артем Бондарчук', 'Артем(Л)': 'Артем Демин', 'Артем(До)': 'Артем Демин',
    'Вова': 'Владимир Туромша', 'Илья': 'Илья Иванов', 'Паша': 'Павел Чекан',
}
# Исправление написания имён из шапки листа.
NAME_FIX = {'Туромша Владимир': 'Владимир Туромша'}

# Строки навыков на листе матча (первая строка = «Связка», и т.д.).
ROWS = {'p': [3, 4, 5, 6], 'a': [12, 13, 14, 15, 16], 'r': [21, 22, 23, 24, 25],
        'df': [28, 29, 30, 31, 32], 'pd': [35, 36, 37], 'b': [42, 43, 44], 's': [47, 48, 49, 50]}
# Какие строки входят в «всего действий» (n). Для атаки не входит «Сохранение».
N_ROWS = {'p': [3, 4, 5, 6], 'a': [12, 13, 14, 15], 'r': [21, 22, 23, 24, 25], 'df': [28, 29, 30, 31, 32],
          'pd': [35, 36, 37], 'b': [42, 43, 44], 's': [47, 48, 49, 50]}
# Итоговые строки книги (для сверки с суммой по партиям).
XL_TOTAL_ROWS = {'p': 7, 'a': 17, 'r': 26, 'df': 33, 'pd': 38, 'b': 45, 's': 51}
TEAM_ROWS = {'op': 58, 'pt': 59, 'er': 60, 'uf': 61}
SETS = 5
STEP = 6  # колонок на игрока: 5 партий + «Итог»

warnings = []
def warn(msg): warnings.append(msg); print('  ПРЕДУПРЕЖДЕНИЕ:', msg)

num = lambda v: v if isinstance(v, (int, float)) else 0

def parse_score(v):
    """Счёт вида «12-14»; Excel мог превратить его в дату или время."""
    if isinstance(v, datetime.datetime): return (v.month, v.day)
    if isinstance(v, datetime.time): return (v.hour, v.minute)
    if isinstance(v, str):
        m = re.match(r'^\s*(\d+)\s*[-:]\s*(\d+)\s*$', v)
        if m: return (int(m.group(1)), int(m.group(2)))
    return None

def player_names(ws):
    names, c = [], 3
    while True:
        v = ws.cell(1, c).value
        if v in (None, ' ', ''): break
        v = str(v).strip(); names.append(NAME_FIX.get(v, v)); c += STEP
    return names

def read_game(ws, ncols):
    g = {'tm': {}, 'pl': []}
    for key, row in TEAM_ROWS.items():
        v = [num(ws.cell(row, 3 + s).value) for s in range(SETS)]
        g['tm'][key] = v + [sum(v)]
    for pi in range(ncols):
        sets = []
        for s in range(SETS):
            col = 3 + STEP * pi + s
            d = {}
            for k, rows in ROWS.items():
                vals = [num(ws.cell(r, col).value) for r in rows]
                d[k] = vals + [sum(num(ws.cell(r, col).value) for r in N_ROWS[k])]
            sets.append(d)
        tot = {k: [round(sum(sd[k][i] for sd in sets), 2) for i in range(len(sets[0][k]))] for k in ROWS}
        g['pl'].append(sets + [tot])
    return g

def read_log(ws, players, label):
    idx = {n: i for i, n in enumerate(players)}
    def pis(text):
        out = []
        for part in [x.strip() for x in text.split('/') if x.strip()]:
            full = NICKNAMES.get(part)
            if full is None or full not in idx:
                warn(f'{label}: в логе неизвестный игрок «{part}» (добавьте его в NICKNAMES)'); out.append(-1)
            else: out.append(idx[full])
        return out
    sets_out, finals = [], []
    for s in range(SETS):
        c = 1 + STEP * s; prev = (0, 0); rl = []
        for r in range(3, 80):
            sc = parse_score(ws.cell(r, c + 2).value)
            if not sc: continue
            pl = (ws.cell(r, c).value or '').strip(); ev = (ws.cell(r, c + 1).value or '').strip()
            oe = (ws.cell(r, c + 3).value or '').strip(); op = (ws.cell(r, c + 4).value or '').strip()
            da, db = sc[0] - prev[0], sc[1] - prev[1]
            if (da, db) not in ((1, 0), (0, 1)):
                warn(f'{label}, сет {s+1}, строка {r}: счёт {prev[0]}:{prev[1]} -> {sc[0]}:{sc[1]} (пропущена или лишняя строка в логе)')
            w = 1 if da > 0 else 0
            kind = ('W' if w else 'F') if pl else ('E' if w else 'O')
            rl.append([w, kind, ev, oe, op, pis(pl)]); prev = sc
        sets_out.append(rl); finals.append(prev[1] if rl else 0)
    return sets_out, finals + [sum(finals)]

def reconcile(game, label):
    ok = True
    for s in range(SETS):
        P = game['tm']['pt'][s]
        if not P: continue
        W = sum(pl[s]['a'][2] + pl[s]['b'][1] + pl[s]['s'][2] for pl in game['pl'])
        E = sum(pl[s]['p'][3] + pl[s]['a'][3] + pl[s]['r'][4] + pl[s]['df'][4] + pl[s]['pd'][2] + pl[s]['b'][2] + pl[s]['s'][3] for pl in game['pl'])
        if abs(W + game['tm']['op'][s] - P) > .01:
            ok = False; warn(f'{label}, сет {s+1}: очки {P:g} ≠ чистые {W:g} + ошибки соперника {game["tm"]["op"][s]:g}')
        if abs(E - game['tm']['er'][s]) > .01:
            ok = False; warn(f'{label}, сет {s+1}: ошибки игроков {E:g} ≠ ошибки команды {game["tm"]["er"][s]:g}')
        rl = game.get('rl')
        if rl is not None:
            if not rl[s]: warn(f'{label}, сет {s+1}: нет лога розыгрышей'); continue
            ours = sum(x[0] for x in rl[s]); F = sum(1 for x in rl[s] if x[1] == 'F')
            Wn = sum(1 for x in rl[s] if x[1] == 'W'); OE = sum(1 for x in rl[s] if x[1] == 'E')
            for nm, a, b in (('очки', ours, P), ('наши ошибки', F, game['tm']['er'][s]),
                             ('чистые очки', Wn, W), ('ошибки соперника', OE, game['tm']['op'][s])):
                if abs(a - b) > .01: warn(f'{label}, сет {s+1}: лог даёт {nm} = {a:g}, таблица = {b:g}')
    return ok

def check_xl_totals(ws, game, label):
    for pi in range(len(game['pl'])):
        col = 3 + STEP * pi + SETS
        for k, row in XL_TOTAL_ROWS.items():
            xl = num(ws.cell(row, col).value); mine = game['pl'][pi][SETS][k][-1]
            if abs(xl - mine) > .01 and k != 'p':
                warn(f'{label}: «Итог» в книге ≠ сумме партий (навык {k}, игрок {pi+1}): {xl:g} vs {mine:g}')

def fetch_xlsx(src):
    """Локальный путь или ссылка (https/file). Для ссылки скачивает книгу во временный файл."""
    if not re.match(r'^(https?|file)://', src): return src
    req = urllib.request.Request(src, headers={'User-Agent': 'Mozilla/5.0 (stats-updater)'})
    try:
        with urllib.request.urlopen(req, timeout=60) as r: body = r.read()
    except Exception as e:
        sys.exit(f'Не удалось скачать таблицу: {e}')
    if body[:2] != b'PK':
        sys.exit('По ссылке пришёл не xlsx (возможно, закрыт доступ по ссылке или отключено скачивание для читателей).')
    fd, path = tempfile.mkstemp(suffix='.xlsx'); os.write(fd, body); os.close(fd)
    return path

def sheet_no(name):
    m = re.search(r'№\s*(\d+)', name)
    return int(m.group(1)) if m else 1   # у первой «Динамики» номера нет

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('xlsx'); ap.add_argument('html')
    ap.add_argument('-o', '--out'); ap.add_argument('--check', action='store_true')
    a = ap.parse_args()

    wb = openpyxl.load_workbook(fetch_xlsx(a.xlsx), data_only=True)
    game_sheets = [n for n in wb.sheetnames if re.search(r'Игра\s*№', n) and 'Динамика' not in n]
    dyn_sheets = [n for n in wb.sheetnames if 'Динамика' in n]
    if not game_sheets: sys.exit('Не найдено ни одного листа матча («Игра №…»)')
    players = player_names(wb[game_sheets[0]])
    print(f'Игроки ({len(players)}): ' + ', '.join(players))

    games = []
    for gi, name in enumerate(game_sheets):
        ws = wb[name]
        if str(ws['A3'].value).strip() != 'Связка': warn(f'{name}: раскладка листа изменилась (в A3 не «Связка») — проверьте номера строк в ROWS')
        g = read_game(ws, len(players))
        if sum(g['tm']['pt'][:SETS]) == 0: continue  # матч ещё не заполнен
        m = re.match(r'\s*(\d+\.\d+)\s+Игра.*?\(vs\s*(.*?)\)?\s*$', name)
        g['dt'] = m.group(1) if m else name.strip(); g['opp'] = m.group(2).strip() if m else ''
        label = f'Матч {g["dt"]} vs {g["opp"]}'
        print(f'\n{label}: очки {g["tm"]["pt"][:SETS]}')
        dyn = next((d for d in dyn_sheets if sheet_no(d) == sheet_no(name)), None)
        if dyn:
            rl, q = read_log(wb[dyn], players, label)
            g['rl'] = rl; g['tm']['q'] = q
        else:
            warn(f'{label}: нет листа «Динамика» — ход партий для этого матча будет недоступен')
        if reconcile(g, label): print('  очки, ошибки и чистые очки сходятся во всех партиях')
        check_xl_totals(ws, g, label)
        games.append({k: g[k] for k in ('dt', 'opp', 'tm', 'pl') + (('rl',) if 'rl' in g else ())})

    data = {'players': players, 'games': games}
    print(f'\nИтого матчей: {len(games)}, предупреждений: {len(warnings)}')
    if a.check: return

    html = open(a.html, encoding='utf-8').read()
    m = re.search(r'(const DATA =\s*)\{"players".*?\};', html, re.S)
    if not m: sys.exit('В HTML не найден блок «const DATA = {…};»')
    new = m.group(1) + json.dumps(data, ensure_ascii=False, separators=(',', ':')) + ';'
    out = html[:m.start()] + new + html[m.end():]
    target = a.out or a.html
    if not a.out: shutil.copyfile(a.html, a.html + '.bak'); print(f'Резервная копия: {a.html}.bak')
    open(target, 'w', encoding='utf-8').write(out)
    print(f'Страница обновлена: {target}')

if __name__ == '__main__':
    main()
